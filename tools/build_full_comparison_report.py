"""Build the Chinese comparison DOCX from versioned, full-precision CSVs.

Run with the existing python-docx environment; no detector or GPU is required.
"""

import csv
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results/paper_comparison"
METHODS = [
    ("A", "Source Only", "无", "无目标域适配；同一 source checkpoint"),
    ("B", "ST / Mean Teacher", "无", "CGA off；VLST off；伪框回归关闭，不等同于 SF-UT"),
    ("C", "Original CLIP-CGA", "CLIP RN50x64", "vanilla CLIP；legacy CGA；无 SARCLIP / LoRA"),
    ("D", "SARCLIP-CGA", "base SARCLIP", "冻结 ViT-B-32；veto_soft；VLST off；无 LoRA"),
    ("E", "ST + VLST", "base SARCLIP", "CGA off；冻结 VLST 编码器；教师伪标签在线原型"),
    ("F", "SARCLIP-CGA + VLST", "base SARCLIP ×2", "D 的 CGA + E 的 VLST；两个独立冻结编码器"),
]


def rows(name):
    with (RESULTS / name).open(newline="", encoding="utf-8") as stream:
        return list(csv.DictReader(stream))


def paragraph(doc, text):
    doc.add_paragraph(text)


def table(doc, headers, values):
    tab = doc.add_table(rows=1, cols=len(headers))
    tab.style = "Light Shading Accent 1"
    for cell, text in zip(tab.rows[0].cells, headers):
        cell.text = text
    repeat = OxmlElement("w:tblHeader")
    tab.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values_row in values:
        for cell, value in zip(tab.add_row().cells, values_row):
            cell.text = str(value)
    for row in tab.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(3)
                for run in para.runs:
                    run.font.size = Pt(9)
    return tab


def matrix(doc, raw, title):
    doc.add_heading(title, level=2)
    corruptions = list(dict.fromkeys(r["corruption"] for r in raw
                                    if r["corruption"] not in ("clean", "clean_val")))
    lookup = {(r["corruption"], r["method"]): float(r["mAP50"]) for r in raw}
    table(doc, ["Corruption", *"ABCDEF"], [
        [c, *(f"{lookup[c, m]:.4f}" for m in "ABCDEF")] for c in corruptions
    ])


def main():
    rsar = rows("raw_results.csv")
    dior = rows("dior_raw_results.csv")
    rsar_metrics = rows("method_mpc.csv")
    dior_metrics = rows("dior_summary_metrics.csv")
    doc = Document()
    section = doc.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21)
    section.top_margin = section.bottom_margin = Cm(1.7)
    section.left_margin = section.right_margin = Cm(1.6)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.paragraph_format.space_after = Pt(6)
    for name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.color.rgb = RGBColor.from_string("203864")
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("IRAOD · seed 42 · ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = "IRAOD 严格 Source-Free 旋转检测完整比较报告"
    doc.core_properties.subject = "RSAR 7×6 与 DIOR-R 3×6；单种子 final-EMA"
    doc.core_properties.author = "IRAOD"

    doc.add_heading("IRAOD 严格 Source-Free\n旋转目标检测比较报告", 0)
    paragraph(doc, "交付日期：2026-09-06｜集成分支：exp/strict-af-integration｜seed=42")
    paragraph(doc, "完成范围：RSAR 42/42 与 DIOR-R 18/18 个 A–F 主表单元格。"
              "A 为固定 Source Only，B–F 为最后一次迭代的 EMA Teacher，不是 60 个适配模型。"
              "DIOR source clean val/test 分别为 0.410 / 0.295；18 个腐蚀单元格加这两项共 20 项评估。"
              "本次只集成现有产物，不训练、不补跑、不选择 checkpoint。")
    paragraph(doc, "结论：RSAR 上 C 的 mPC 最高（0.3336）；DIOR-R 上 E 最高（0.2757）。"
              "完整组合 F 在两套数据的平均适配增益均为负，现有证据不支持“组合优于各单模块”。"
              "所有结果只有一个种子，std、置信区间和显著性检验均为 N/A，不作 SOTA 声明。")
    paragraph(doc, "数值来源：RSAR 固定沿用提交 b8527464c0201efb195d1f311f82bd86fb4976ca；"
              "DIOR 来自 gpu-67 上 comparison/deliverables 的原始 CSV，运行代码为 "
              "0f98a48b539f9260055fc305784ffbc924f50451。完整二进制仅由远端 manifest 定位。")

    doc.add_heading("1. 协议、公平性与泄漏约束", 1)
    paragraph(doc, "评测为 OBB 旋转框 VOC 风格 AP50 / mAP50（IoU=0.5，DOTA le90），"
              "不是 COCO AP50:95，也不与 HBB、source-available UDA 或 Target-supervised oracle 混排。"
              "A–F 在各数据集内部共享 Oriented R-CNN + OrthoNet-50–FPN 与同一 source；"
              "两数据集使用各自 clean train 训练的源模型，不能把二者当成同一 checkpoint。")
    table(doc, ["数据集", "Source clean TEST", "无标注适配", "预算 / 角色"], [
        ["RSAR（6 类）", "0.5313997865", "每种 corruption 的 val 图像 8467 张",
         "1 epoch / 265 iter；B–F final EMA"],
        ["DIOR-R（20 类）", "0.2948998511", "每种 corruption 的 val 图像 5863 张",
         "1 epoch / 185 iter；B–F iter_185_ema.pth"],
    ])
    paragraph(doc, "共同适配设置：seed 42、deterministic、有效 batch=32、LR=0.02；"
              "B/C/D 为 1×32，E/F 为 2×16（OOM 后保持全局 batch 的拓扑调整）。"
              "共享基础 ST 配方、增广、伪标签基础阈值、rotated NMS 和评测器，仅改变规定的 CGA/VLST 模块。"
              "拓扑相同预算不代表浮点轨迹完全一致。")
    paragraph(doc, "严格边界：weight_l=0、weight_u=1；StrictSourceFreeDOTADataset 使用 image-only val，"
              "返回空 gt_bboxes/gt_labels；--no-validate。GT 不进入训练、伪标签、阈值选择、"
              "prototype 初始化、early stopping 或 checkpoint 选择；测试标注只用于离线评测/展示。"
              "use_bbox_reg=False 是共同 SFOD 配方，并非遗漏监督分支。")
    paragraph(doc, "C 使用冻结 vanilla CLIP；D/E/F 使用冻结 base SARCLIP，禁用目标标注 LoRA。"
              "启动脚本清理 CGA_*、SARCLIP_*、VLST_* 后注入各方法环境；VLST 原型只由教师伪标签在线更新。"
              "预训练基础模型的数据溯源限于现有权重与审计，不能宣称另行核验了全部预训练样本。")
    paragraph(doc, "DIOR source 在 clean train 训练，clean val 0.4098773599 仅作源模型诊断；"
              "所有 DIOR rPC 分母使用 clean TEST 0.2948998511，绝不使用较高的 val 分数。")
    paragraph(doc, "审计入口：experiments/comparison/leakage_audit.md、protocol_audit.md；"
              "results/paper_comparison/dior_audit.md、dior_checkpoint_manifest.json；"
              "源配置与启动信息保存在 experiments/comparison/dior_source/。")

    doc.add_heading("2. A–F 定义与公平性分组", 1)
    table(doc, ["ID", "方法", "VLM", "唯一模块差异 / 状态"], METHODS)
    paragraph(doc, "上述主表均为 Source-Free / OBB / 不使用目标域标签；"
              "A 没有适配阶段，B–F 使用最后 EMA。B 关闭 RPN/ROI 伪框回归，"
              "不等同于保留这两项回归损失的 SF-UT，不能用 B 数值填充 SF-UT。"
              "Target-supervised oracle 未纳入本次交付，数值 N/A。")

    doc.add_heading("3. 完整逐腐蚀主结果", 1)
    matrix(doc, rsar, "3.1 RSAR：7 corruption × A–F（42/42）")
    matrix(doc, dior, "3.2 DIOR-R：3 corruption × A–F（18/18）")
    paragraph(doc, "表内 mAP50 保留四位小数；精确计算来自 raw CSV 的浮点值，不从已四舍五入的表格反算。"
              "DIOR source clean val/test 三位小数为 0.410 / 0.295；B–F 未提供独立 clean 适配评估，"
              "其 clean 单元格为 N/A，不能把 A clean 分数填作适配模型实测值。")

    doc.add_heading("4. mPC / rPC / 平均适配增益", 1)
    paragraph(doc, "mPC = 全部 corruption mAP50 的算术平均；"
              "rPC = mPC / A_clean_TEST × 100%；"
              "mean Δada = mean(mAP_method − mAP_A)，在相同 corruption 上配对。"
              "Δ 的 0.01 相当于 1 个百分点，不能与 1% 相对提升混淆。")
    for dataset, metrics, rpc_key in (
        ("RSAR", rsar_metrics, "rPC"), ("DIOR-R", dior_metrics, "rPC_percent")
    ):
        doc.add_heading(dataset, 2)
        table(doc, ["方法", "mPC", "rPC (%)", "mean Δada", "相对 A mPC (%)", "std"], [
            [r["method"], f'{float(r["mPC"]):.4f}', f'{float(r[rpc_key]):.2f}',
             f'{float(r["mean_delta_ada"]):+.4f}',
             f'{100 * float(r["mean_delta_ada"]) / float(metrics[0]["mPC"]):+.2f}',
             "N/A（单种子）"] for r in metrics
        ])
    paragraph(doc, "DIOR-R mPC 按 A/B/C/D/E/F 的三位小数为 "
              "0.271 / 0.275 / 0.275 / 0.268 / 0.276 / 0.271。"
              "B=0.2754657219、C=0.2754480441、E=0.2756796877 的差距很小；"
              "四舍五入并列不等于统计等价，单次排序也不等于稳定优势。")

    doc.add_heading("5. Recovery 不稳定性与统计解释", 1)
    paragraph(doc, "recovery = (mAP_method − mAP_A) / (A_clean − A_corruption)，"
              "仅在分母为正时定义；否则为 N/A。mean recovery 是逐腐蚀比值的平均，"
              "不是 mean Δada 除以平均性能下降。它可能与 mean Δada 的符号不同，不用于主排名。")
    table(doc, ["方法", "RSAR mean recovery", "DIOR mean recovery"], [
        [r["method"], f'{float(r["mean_recovery"]):+.3f}',
         f'{float(d["mean_recovery"]):+.3f}']
        for r, d in zip(rsar_metrics, dior_metrics)
    ])
    paragraph(doc, "DIOR cloudy 分母仅 0.004631579，contrast 为 0.013296723，均被远端标记 unstable；"
              "cloudy 的 E recovery 约 1.577 只是适配后略超 source clean 分数，不代表稳健的“157.7% 全面恢复”。"
              "RSAR gaussian_white_noise 的分母仅 0.003262639，point_target 0.015559673，"
              "chaff 0.059468895；小幅负迁移会放大，使 C/E 的 mean recovery 为负而 mean Δada 为正。")
    paragraph(doc, "单种子 seed=42：跨 corruption 的变化不是跨 seed 标准差；"
              "不以 corruption 充当独立重复实验，不报告虚构 std、p 值或显著性。"
              "如需后续统计结论，应先冻结协议再进行独立多种子重复；本次未启动新实验。")

    doc.add_heading("6. Per-class 与数值追溯", 1)
    per_class = rows("dior_per_class.csv")
    coverage = Counter((r["corruption"], r["method"]) for r in per_class)
    table(doc, ["索引文件（results/paper_comparison/）", "行数", "覆盖 / 精度"], [
        ["dior_raw_results.csv", len(dior), "18 corruption + source clean val/test"],
        ["dior_per_corruption.csv", 18, "逐格 Δada / recovery / 稳定性标记"],
        ["dior_per_class.csv", len(per_class), "20 个评估 × 20 类 AP50；类 AP 为日志三位小数"],
        ["dior_evaluation_manifest.csv", 20, "原始 eval JSON + config + checkpoint 绝对路径"],
        ["per_class_summary.csv（RSAR）", len(rows("per_class_summary.csv")),
         "既有 78 行；不是 RSAR 全矩阵 per-class 完整导出"],
    ])
    paragraph(doc, "DIOR per-class 分组：" + "；".join(
        f"{c}/{m}: {n} 类" for (c, m), n in coverage.items()) + "。")
    paragraph(doc, "类别顺序见 dior_checkpoint_manifest.json（expressway 两类在 dam 之前，统一小写）。"
              "per-class 的 AP 已四舍五入到三位小数，其均值可能与 full-precision mAP 略有差异，"
              "不得用此差异改写 raw_results。原始日志、预测与权重仍留在远端；"
              "manifest 记录的 eval JSON 和 checkpoint 才是本地可追溯入口。")

    doc.add_heading("7. 同图可视化与 RoI 文件定位", 1)
    paragraph(doc, "DIOR 固定使用 ImageSets/test.txt 的前 16 个 ID（11726–11741），"
              "三个 corruption、六种方法使用同一集合；不依据任何方法的结果挑图。"
              "共 3×6×16=288 张检测图，show-score-thr=0.3；"
              "共享源评测配置的 score_thr=0.05、rotated NMS IoU=0.1，展示阈值不等于评测阈值。"
              "为控制仓库体积，本 DOCX 不嵌入图片，也未据此声称具体图像上的定性优势。")
    table(doc, ["本地 manifest", "记录数", "用途"], [
        ["dior_visualization_manifest.csv", 51, "远端原始目录汇总：18 vis + 33 RoI"],
        ["dior_visualization_files.csv", 288, "逐图 image / method / config / checkpoint / show_dir"],
        ["dior_roi_manifest.csv", 33, "3×(A source + B–F EMA/Student)，每项 16 图"],
        ["dior_visualization_selection.json", "16 IDs", "固定 test 前 16 图，无挑图"],
    ])
    paragraph(doc, "远端根目录（gpu-67）："
              "/mnt/shared/zechuan/iraod_artifacts/comparison/deliverables/results/paper_comparison/")
    paragraph(doc, "检测图：visualizations/DIOR/<corruption>/<A-F>/<image_id>.jpg。"
              "例：visualizations/DIOR/brightness/A/11726.jpg；"
              "同图 B–F 只替换方法目录。")
    paragraph(doc, "RoI：roi_features/DIOR/<corruption>/<A-F>/<ema|student>/<image_id>.npy；"
              "同目录 index.json 保存 image_id、n_roi、preds[class_id, score, rotated_box]。"
              "A 的 ema 目录实际保存 source 特征，无独立 Student；"
              "B–F 同时保留最后 EMA 与 Student，共 33 目录 / 528 个特征文件。"
              "所有特征提取点均为 roi_head.bbox_head.fc_cls 输入。")
    paragraph(doc, "重要限制：当前导出是 pre-NMS RoI feature 行 + post-NMS 检测 preds，"
              "没有逐行匹配索引，不能把 preds 直接当作每个 feature 的标签/框。"
              "33 个目录存在不等于“带逐实例标签的 t-SNE 已就绪”；"
              "可做无标签特征分析，按预测类别着色的实例级分析需先补齐 proposal/NMS 行映射。"
              "这属于原始交付字段未完全满足的限制，本次没有修改或重跑特征导出。")
    paragraph(doc, "RSAR：已有 run_manifest 记录 chaff 32 图的可视化及部分 RoI 状态，"
              "但本地 visualization_manifest.csv 仍是 not_started 占位行，"
              "本次未扩展其二进制或逐图索引；不能据 42/42 数值完成推断七种干扰的定性产物全覆盖。")

    doc.add_heading("8. 外部基线：N/A 的具体原因", 1)
    paragraph(doc, "N/A 表示当前冻结框架没有经过忠实移植、验证及同协议评估的数值，"
              "不是方法失效，也不是声称理论上无法移植。作者仓库地址来自原始计划；"
              "除已有远端审计外，本次未重新审读这些官方实现，不能编造实现细节。")
    table(doc, ["方法 / 官方仓库", "状态", "忠实移植边界"], [
        ["IRG-SFDA\nvibashan/irg-sfda", "N/A",
         "缺 OBB 实例关系图模块与 rotated IoU/proposal 对应实现；不能把 HBB 图机制直接接入共享源并宣称已复现。"],
        ["LPLD\nCV-Det/LPLD", "N/A",
         "冻结树与远端无对应代码/config/checkpoint；伪标签蒸馏及 proposal 对应的 rotated 匹配未移植、未验证，不能以 B 改名替代。"],
        ["DRU\nlbktrinh/DRU", "N/A",
         "0f98a48 中无实现；尚未核验其核心机制能否在固定 source 和预算下独立于原 detector 移植，不能用近似替身填表。"],
        ["AASFOD / A²SFOD\nChuQiaosong/AASFOD", "N/A",
         "无可执行 OBB port；模块、assignment/IoU/NMS 的旋转语义及无目标 GT 配方未验证，现有材料不足以宣称忠实重现。"],
        ["SF-YOLO\nvs-cv/sf-yolo", "N/A",
         "原 YOLO detector 不同于共享 Oriented R-CNN；直接使用 YOLO 或另训 source 会破坏主表公平性，独立核心机制 OBB 移植尚未完成。"],
        ["SF-UT / Simple-SFOD\nEPFL-IMOS/simple-SFOD", "需独立实验",
         "SF-UT 保留 RPN/ROI 伪框回归；B 将两者置零。AdaBN+Fixed SF-FixMatch 又是不同策略，不能统一并入 B。"],
    ])

    doc.add_heading("9. Recovery commands 与复现材料", 1)
    paragraph(doc, "从本地仓库执行以下命令连接 gpu-67，使用现有远端环境和绝对路径。"
              "这些是恢复入口，不代表本次执行过训练。"
              "原脚本的 skip 条件并非通用断点恢复：训练匹配任意 iter_*_ema.pth，"
              "评估检查 eval_exit=0；中断产物须人工确认 iter_185_ema.pth 与成功状态并隔离残缺输出，"
              "不能把中间 checkpoint 当作 final EMA。缺数据、权重、原环境或脚本时不能保证恢复。")
    paragraph(doc, "单格（示例：B / brightness）：\n"
              "ssh gpu-67 'bash -s -- cell B brightness 4' "
              "< results/paper_comparison/dior_recovery_commands.sh")
    paragraph(doc, "完整矩阵（顺序运行；E/F 使用指定双卡）：\n"
              "ssh gpu-67 'bash -s -- matrix 4 4,5' "
              "< results/paper_comparison/dior_recovery_commands.sh")
    paragraph(doc, "仅 source clean val/test：\n"
              "ssh gpu-67 'bash -s -- source 4' "
              "< results/paper_comparison/dior_recovery_commands.sh")
    paragraph(doc, "按需重新生成可视化 / RoI 的原始入口："
              "远端 comparison/dior/queue_0f98a48/run_vis.sh <gpu> <method> <corruption>；"
              "run_roi.sh <gpu> <method> <corruption> <ema|student>。"
              "原始 run_roi 只判断 index.json 是否存在；残缺输出应由人工确认并隔离后恢复，"
              "不得把存在文件当作完整成功。")
    paragraph(doc, "小型脚本、DIOR overlay 与 source resolved config 已存入 "
              "experiments/comparison/dior_recovery/ 和 dior_source/；"
              "它们是保留原绝对路径的执行快照，不是可在本地 CPU 环境直接运行的训练入口。"
              "RSAR 原有恢复入口保持 experiments/comparison/reproduction_commands.sh。")

    doc.add_heading("10. 主要结论、限制与交付边界", 1)
    paragraph(doc, "① 观测：RSAR C/E 较 A 分别增加 0.0165 / 0.0142 mPC；"
              "提升主要来自强干扰，而 chaff/point_target/gaussian 上出现负迁移。"
              "解释：语义模块收益具有干扰依赖性，不能笼统声称全场景提升。")
    paragraph(doc, "② 观测：DIOR B/C/E 为约 +0.0043 / +0.0043 / +0.0045，"
              "D 为 −0.0031，F 为 −0.0005。"
              "解释：冻结 SARCLIP-CGA 及其组合在光学 DIOR 上没有平均增益；"
              "模态错配或模块冲突仅是待检验假设，不能由聚合值确定因果。")
    paragraph(doc, "③ 观测：F 在 RSAR 七种干扰均低于 A，平均 −0.0197，DIOR 也未超过 B/C/E。"
              "含义：论文应如实呈现负结果与消融，不通过挑图、挑 checkpoint 或换分母美化完整方法。")
    paragraph(doc, "完成的是既有数值与可追溯小型产物的集成报告，不是宣称原计划所有扩展交付均完成。"
              "限制包括：单种子、仅三种已存在 DIOR corruption、RSAR per-class/定性覆盖不完整、"
              "RoI 无 pre/post-NMS 逐行映射、外部五类方法 N/A、oracle 未跑；"
              "未重新核验全部基础模型预训练数据，也未对远端全部图片逐张做人工视觉判断。"
              "DIOR 可视化与 mAP 是独立 test.py 调用：前者固定 16 图，后者完整 TEST，不能混淆。")
    paragraph(doc, "仓库只保存小型 CSV / LaTeX / Markdown / JSON / YAML / recovery scripts 与本 DOCX。"
              "不提交 checkpoint、数据、全量图片/特征或日志；用户主 checkout 的“对比实验的plan.docx”"
              "保持原样且不提交。原始计划字段依据用户指定的 plan.original.docx 逐项对照。")
    output = RESULTS / "IRAOD_full_comparison_report_cn.docx"
    doc.save(output)
    print(f"Generated {output.relative_to(ROOT)} ({output.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
