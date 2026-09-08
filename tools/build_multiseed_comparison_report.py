"""Render the supplied multi-seed evidence JSON as a Chinese DOCX (CPU only)."""

import argparse
import json
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from tools.build_full_comparison_report import table


def number(value):
    return "N/A" if value is None else f"{value:.6g}"


def grouped_tables(doc, report, records, headers, format_row):
    """Never mix the extended-budget or target-supervised appendix into strict tables."""
    if "comparison_groups" not in report:
        table(doc, headers, [format_row(r) for r in records])
        return
    for group, methods in report["comparison_groups"].items():
        doc.add_heading(group, 2)
        table(doc, headers, [format_row(r) for r in records if r["method"] in ["A", *methods]])


def render(report_path, output):
    path = Path(report_path)
    report = json.loads(path.read_text())
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "SimSun")
    doc.add_heading("IRAOD 多种子结果与证据覆盖报告", 0)
    doc.add_paragraph(f"结果状态：{report['status']}。本文件是消费者产物，"
                      "报告生成成功不等于实验或全部八项任务完成。")
    doc.add_paragraph("输入清单：" + report["input_manifest"])
    doc.add_paragraph("消费者提交：" + report["code_commit"])
    if report.get("findings_cn"):
        doc.add_heading("主要观察与统计限制", 1)
        for text in report["findings_cn"]:
            doc.add_paragraph(text)
    corrections = report.get("producer_metadata_corrections", [])
    if corrections:
        doc.add_heading("Source producer 元数据更正审计", 1)
        doc.add_paragraph("以下仅更正报告中的 source producer 解释：原 sidecar 保留，"
                          "权重、预测、实际图像 ID 和 mAP 均不改变，不因此重训或重推理。"
                          "recorded 与 actual 明确分列；B–F 的适配训练 producer 不套用 source producer。")
        table(doc, ["数据集/域", "sidecar recorded", "source actual", "原始依据"], [
            [f"{r['dataset']}/{r['domain']}", str(r["recorded"]), r["actual"], r["source_record"]]
            for r in corrections])
    doc.add_heading("1. 统计单位与适用范围", 1)
    doc.add_paragraph(
        "实验单位为适配种子 42、43、44。先在每个种子内对全部腐蚀域求 mPC，"
        "并在同域求相对固定 A 的差值后平均，再跨种子计算均值与样本标准差（ddof=1）。"
        "A 仅一次固定 source 测量，不复制成三次独立测量；所有区间只反映条件于固定 source 的适配随机性。"
        "不把类别、图像或腐蚀域当作独立种子，也不要求 B–D 单卡和 E/F 双卡具有相同 GPU UUID。")
    doc.add_paragraph(
        "n=3，检验低功效。95% CI 为 100000 次 seed-block 百分位 bootstrap（随机种子 42），"
        "小样本下不稳定。双侧 t 检验施于三个种子级配对差值，依赖差值正态假设；"
        "零方差时 t 检验未定义。精确 sign-flip 枚举 8 种符号，其最小双侧 p=0.25。"
        + ("Holm 按清单预先声明的方法、数据集、角色、指标及比较组校正；家族不完整不输出调整后 p 值。"
           if "methods" in report else
           "Holm 分别用于各数据集/角色/指标的五个 B–F 比较，家族不完整时不输出调整后 p 值。"))
    if "methods" in report:
        doc.add_paragraph("声明方法：" + ", ".join(report["methods"]))
        doc.add_paragraph(
            "common_one_epoch 为共同一轮预算；AASFOD 使用已披露的预算缩放移植。"
            "SFYOLO 为 extended_two_epoch_TAM，两轮最终标签 RSAR531 / DIOR369，"
            "不与一轮方法合并排名或统计。LoRA-CGA / LoRA-CGA+VLST 如被声明，"
            "仅在 target_supervised_appendix 中展示，绝不计入 strict main。"
            "各表与图按清单顺序而非结果排序；固定 A 在各组仅作同一参考，并非新增测量。")
    doc.add_paragraph(
        "历史 rPC 始终为 100 × mPC / 固定 A clean TEST；"
        "method_clean_normalized_percent 另列，分母为同方法同种子的独立 clean 适配实测值。"
        "缺域或缺种子时不生成完整多种子均值、标准差或 CI，不据少量已完成结果排名。")
    doc.add_paragraph(
        "Recovery ratio = (方法 corruption AP50 − A corruption AP50) / "
        "(A clean TEST AP50 − A corruption AP50)。分母非正时未定义，接近零时不稳定；"
        "不截断负值或大于 1 的值，不据此排名，逐种子逐域数值见 recovery.csv。")
    if report.get("method_matrix"):
        doc.add_heading("公平性与方法边界", 1)
        table(doc, ["方法", "公平组", "框", "VLM", "目标标签", "状态"], [
            [r["id"], r["fairness_group"], r["box"], r["vlm"], r["target_labels"], r["status"]]
            for r in report["method_matrix"]])
        doc.add_paragraph("本报告的 Source-Free A–F 为同一比较组。B/D/E/F 为已有组件组合比较，"
                          "不冒充额外消融训练。未列出分数的扩展方法与 target-supervised oracle"
                          "不在本报告的结果范围内；这不表示它们未获批准或已完成。")
    doc.add_heading("2. 量化与逐实例覆盖", 1)
    cov = report["qualitative_coverage"]
    coverage_rows = [
        ["量化单元", f"{report['quantitative_complete_cells']} / {report['quantitative_expected_cells']}",
         "全 TEST predictions；需实际有序 image-ID 证据"],
    ]
    if report.get("quantitative_only"):
        doc.add_paragraph(report["collection"]["export_scope"])
        doc.add_paragraph(
            "本次只收集声明角色的全 TEST 量化结果；不把既有 core 的定性覆盖分母"
            "当作扩展任务范围。新增 RoI、固定可视化与联合嵌入仍由已批准扩展清单跟踪。")
    else:
        expected_vis = cov["vis_expected_image_roles"] if "methods" in report else 3520
        expected_embeddings = cov["embeddings_expected"] if "methods" in report else 24
        coverage_rows.extend([
        ["RoI", f"{cov['roi_complete']} / {cov['roi_expected_image_roles']}", "全 TEST 对齐逐实例"],
        ["可视化", f"{cov['vis_complete']} / {expected_vis}", "固定子集，展示阈值 0.3"],
        ["joint embedding", f"{cov['embeddings_complete']} / {expected_embeddings}",
         "采样分析；不是全量预测实例"],
        ])
    table(doc, ["证据类型", "已完成 / 预期", "范围"], coverage_rows)
    doc.add_paragraph(
        "predictions.pkl 与原生 image-ID 顺序是量化证据，不包含 fc_cls 特征，"
        "不能替代全 TEST aligned RoI；定性导出未在本次执行。"
        if report.get("quantitative_only") else
        "全 TEST 图片的 aligned RoI 保存全部 post-NMS 检测，不受展示阈值 0.3 二次截断；"
        "检测器自身的 score threshold、NMS 和 max_per_img 仍生效。"
        "可视化仅使用冻结的 RSAR32/DIOR16 子集；不影响全量 RoI 完成。"
        "旧 v2 子集不计入 full-test 完成。predictions.pkl 不含 fc_cls 特征，不能替代 aligned RoI；"
        "pred_count 相符也不能替代图像顺序证据。")
    grouped_tables(doc, report, report["per_domain"],
                   ["数据集/域", "方法/角色", "n", "均值", "样本 std", "状态"], lambda r:
        [f"{r['dataset']}/{r['domain']}", f"{r['method']}/{r['role']}", r["n"],
         number(r["mean"]), number(r["sample_std"]), r["status"]]
    )
    doc.add_heading("3. 多种子汇总", 1)
    grouped_tables(doc, report, report["summary"],
                   ["数据集/角色", "方法", "指标", "n", "均值", "样本 std", "95% CI", "状态"], lambda r:
        [f"{r['dataset']}/{r['role']}", r["method"], r["metric"], r["n"],
         number(r["mean"]), number(r["sample_std"]),
         f"[{number(r['bootstrap_ci95_low'])}, {number(r['bootstrap_ci95_high'])}]",
         r["status"]]
    )
    for figure in report["figures"]:
        if figure.endswith(".png"):
            doc.add_paragraph(figure + "：仅展示完整种子块，误差棒为样本标准差；A 无重复测量误差棒。")
            doc.add_picture(str(path.parent / figure), width=Inches(5.5))
    doc.add_heading("4. 配对统计", 1)
    grouped_tables(doc, report, report["paired_statistics"],
                   ["数据集/角色", "方法", "指标", "n", "t p", "sign-flip p", "Holm t / sign", "状态"], lambda r:
        [f"{r['dataset']}/{r['role']}", r["method"], r["metric"], r["n"],
         number(r["t_p"]), number(r["sign_flip_p"]),
         f"{number(r['t_p_holm'])} / {number(r['sign_flip_p_holm'])}", r["status"]]
    )
    doc.add_heading("5. 原始值、精度与未完成项", 1)
    doc.add_paragraph(
        "raw_results.csv 保留完整 metric.mAP，不使用已舍入的 metric.AP50。"
        "historical/ 逐字节保留输入的历史 seed42 CSV；原始 seed42 mAP 锁定，"
        "重评差异另存 eval_mAP50 并标记 metric_conflict，不静默改写旧值。"
        "per_class.csv 来自 class_ap.txt，通常仅三位小数，不反算总 mAP。"
        "prediction_image_coverage.csv、roi_vis_coverage.csv 和 embedding_index.csv 是三种不同范围的证据。")
    missing = [r for r in report["raw_results"] if r["status"] != "complete"]
    table(doc, ["数据集/域", "方法/角色/种子", "状态", "原因"], [
        [f"{r['dataset']}/{r['domain']}", f"{r['method']}/{r['role']}/{r['seed']}",
         r["status"], "; ".join(r["problems"])] for r in missing])
    doc.add_paragraph(
        "全部原始行与每种子聚合见 raw_results.csv、per_seed.csv；"
        "本报告不声称完成其他基线移植或未在输入证据中完成的项目。")
    doc.save(output)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--report", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    render(args.report, args.out)


if __name__ == "__main__":
    main()
